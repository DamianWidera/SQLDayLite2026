#!/usr/bin/env python3
"""Build the workshop deck "Kupiłeś Fabric. I co teraz?" from the SQLDay Lite 2026 template.

Usage:
    python build_deck.py [output.pptx]

Reads  slides/SQLDay_Lite_2026.pptx (template, never modified) and ../assets/*.
Writes slides/Kupiles-Fabric-I-co-teraz.pptx (or the path given as the first argument).
Requires python-pptx >= 1.0.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
TEMPLATE = HERE / "SQLDay_Lite_2026.pptx"
OUTPUT = HERE / "Kupiles-Fabric-I-co-teraz.pptx"
ARCH = ROOT / "assets" / "architecture"
QR = ROOT / "assets" / "qr-ankieta.png"
LOGO_DC = ROOT / "assets" / "logodc.png"

# Layout indexes in the template
L_TITLE = 0          # light
L_LIGHT = 1          # Title and Content, light
L_DARK_WARM = 2      # 1_Title and Content, dark with warm gradient
L_DARK = 3           # 2_Title and Content, dark slate
L_DARK2 = 4          # 3_Title and Content, dark slate
L_DARK_BLUE = 6      # 5_Title and Content, dark with blue glow
L_BLUE = 7           # 6_Title and Content, blue
L_SECTION = 8        # Section Header, light
L_TWO = 9            # Two Content, light
L_COMPARE = 10       # Comparison, light
L_TITLE_ONLY = 11    # Title Only, light

NAVY = RGBColor(0x0D, 0x14, 0x33)
ORANGE = RGBColor(0xFF, 0x7A, 0x1A)
LORANGE = RGBColor(0xFF, 0xB3, 0x47)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xF1, 0xF2, 0xF7)      # light card fill on light layouts
SILVER = RGBColor(0xC3, 0xC9, 0xD6)
BRONZE = RGBColor(0xC9, 0x7A, 0x3A)
GREY = RGBColor(0x5A, 0x60, 0x75)
FONT = "Oswald"                        # the template's master font

# The SQLDay Lite logo is baked into the layout backgrounds, top right.
# Keep x > 10.3 in, y < 1.5 in free.
LEFT = 0.92
WIDTH = 11.5


# ---------------------------------------------------------------- helpers
def _run(par, text, size=None, color=None, bold=None, font=FONT, italic=None):
    r = par.add_run()
    r.text = text
    f = r.font
    if size:
        f.size = Pt(size)
    if color is not None:
        f.color.rgb = color
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if font:
        f.name = font
    r._r.get_or_add_rPr().set("lang", "pl-PL")
    return r


def write(tf, paragraphs, align=None, anchor=None, margins=None, space_after=None,
          line_spacing=None, bullets=None):
    """paragraphs: list of paragraphs, each a list of run tuples (text, size, color, bold)
    or a single run tuple. Replaces all text in the text frame."""
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    if margins is not None:
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    first = True
    for runs in paragraphs:
        if isinstance(runs, tuple):
            runs = [runs]
        par = tf.paragraphs[0] if first else tf.add_paragraph()
        if first:
            for r in list(par.runs):
                r._r.getparent().remove(r._r)
        first = False
        if align is not None:
            par.alignment = align
        if space_after is not None:
            par.space_after = Pt(space_after)
        if line_spacing is not None:
            par.line_spacing = line_spacing
        if bullets is False:
            pPr = par._p.get_or_add_pPr()
            pPr.set("marL", "0")
            pPr.set("indent", "0")
            for child in list(pPr):
                if child.tag.endswith("}buChar") or child.tag.endswith("}buNone"):
                    pPr.remove(child)
            pPr.append(pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone", {}))
        for run in runs:
            _run(par, *run)


def no_autofit(tf):
    bodyPr = tf._txBody.bodyPr
    for child in list(bodyPr):
        if child.tag.endswith("Autofit"):
            bodyPr.remove(child)
    bodyPr.append(bodyPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}noAutofit", {}))


def box(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.08, dash=False,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    return s


def textbox(slide, x, y, w, h, paragraphs, **kw):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    kw.setdefault("margins", (0.05, 0.03, 0.05, 0.03))
    write(t.text_frame, paragraphs, **kw)
    return t


def card(slide, x, y, w, h, heading, body, dark=True, head_size=24, body_size=17,
         accent=ORANGE, number=None):
    """Card with accent bar on top, heading and body text."""
    fill = NAVY if dark else MIST
    head_color = LORANGE if dark else NAVY
    body_color = WHITE if dark else NAVY
    s = box(slide, x, y, w, h, fill=fill, radius=0.06)
    box(slide, x + 0.22, y + 0.2, 0.7, 0.07, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    paragraphs = []
    if number is not None:
        paragraphs.append([(number, head_size, ORANGE, True)])
    paragraphs.append([(heading, head_size, head_color, True)])
    if body:
        paragraphs.append([(body, body_size, body_color, False)])
    textbox(slide, x + 0.12, y + 0.36, w - 0.24, h - 0.45, paragraphs, space_after=6,
            anchor=MSO_ANCHOR.TOP)
    return s


def tbc(slide, x, y, w, h, label, dark=False, size=14):
    """Marked placeholder the speakers still have to fill."""
    color = LORANGE if dark else ORANGE
    s = box(slide, x, y, w, h, fill=None, line=color, line_w=1.5, dash=True, radius=0.1)
    write(s.text_frame, [[(label, size, WHITE if dark else NAVY, False)]],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margins=(0.08, 0.04, 0.08, 0.04))
    return s


def picture(slide, path, x, y, w=None, h=None):
    """Add a picture keeping its aspect ratio (only one of w, h is used)."""
    if w is not None:
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(h))


def place(shape, x, y, w, h):
    """Set all four geometry values (a placeholder inherits nothing once one is set)."""
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def drop_placeholder(slide, idx):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx == idx:
            ph._element.getparent().remove(ph._element)


def set_title(slide, text, size=40, color=None, top=0.4, height=1.1, width=9.3, sub=None,
              sub_color=None, sub_size=22):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(LEFT), Inches(top), Inches(width), Inches(height)
    paragraphs = [[(text, size, color, None, None)]]
    if sub:
        paragraphs.append([(sub, sub_size, sub_color, None, None)])
    write(t.text_frame, paragraphs, anchor=MSO_ANCHOR.MIDDLE)
    no_autofit(t.text_frame)
    return t


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


class Deck:
    def __init__(self):
        self.prs = Presentation(str(TEMPLATE))
        self._original = list(self.prs.slides._sldIdLst)

    def slide(self, layout, title=None, **kw):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[layout])
        if title is not None:
            set_title(s, title, **kw)
        return s

    def save(self, path):
        lst = self.prs.slides._sldIdLst
        for sldId in self._original:           # remove the template's sample slides
            self.prs.part.drop_rel(sldId.rId)
            lst.remove(sldId)
        self.prs.save(str(path))


# ---------------------------------------------------------------- content
EXERCISES = [
    dict(
        img="architektura-cw0.png", big="0", kicker="Start", dur="45 min", start="08:30",
        title="Start i konfiguracja",
        sub="Logowanie, gotowy workspace, Spark pool i pliki do ćwiczeń",
        goal="Zalogować się do Fabric i sprawdzić swój workspace.",
        tasks=[("1", "Zaloguj się w trybie incognito loginem fabric.workshop.sepNNN"),
               ("2", "Poznaj przełącznik Fabric i Power BI oraz Settings"),
               ("3", "Otwórz swój workspace Fabric Workshop September NNN"),
               ("4", "Sprawdź Spark pool: maksymalnie 2 węzły i Runtime 1.3"),
               ("5", "Pobierz pliki do ćwiczeń z repozytorium")],
        outcome="otwarty workspace, mały Spark pool i pliki na dysku.",
        n_div="Pokaż diagram i powiedz, że zaczynamy od lewego dolnego rogu. Każdy loguje się loginem z numerem, ten sam numer ma jego workspace. Trzy capacity, około dziesięć osób na każdej. "
              "Przypomnij o trybie incognito. Sprawdź, czy wszyscy widzą ekran powitalny Fabric.",
        n_tasks="Przejdź kroki ze start.md na żywo. Zatrzymaj się przy kroku 13: sprawdzamy 2 węzły i Runtime 1.3, bo dziesięć osób dzieli capacity. "
                "Nikt nie tworzy nowych workspace. Poczekaj, aż każdy widzi swój workspace.",
    ),
    dict(
        img="architektura-cw1.png", big="1", kicker="Ćwiczenie", dur="80 min", start="09:15",
        title="Ćwiczenie 1",
        sub="Ładowanie danych: Pipeline i Shortcut",
        goal="Załadować surowe dane NYC Taxi do warstwy bronze.",
        tasks=[("1.1", "Pipeline z Copy activity: Azure Blob Storage do Lakehouse"),
               ("1.2", "Poznaj Lakehouse: Files, Tables i format Delta"),
               ("1.3", "Shortcut do ADLS Gen2, potem Load to Tables"),
               ("1.4", "Zarządzanie Spark session")],
        outcome="Lakehouse bronzerawdata z tabelami green_202201_202301 i green202301.",
        n_div="Pokaż na diagramie dwa źródła i dwie drogi do bronze. Pipeline kopiuje dane, Shortcut ich nie kopiuje. "
              "Uprzedź, że to ćwiczenie to głównie klikanie w interfejsie.",
        n_tasks="Zrób zadanie 1.1 razem z grupą, resztę uczestnicy robią sami. Przy 1.2 pokaż folder _delta_log. "
                "Przy 1.4 pokaż Monitoring hub i anulowanie Spark session.",
    ),
    dict(
        img="architektura-cw2.png", big="2", kicker="Ćwiczenie", dur="75 min", start="10:50",
        title="Ćwiczenie 2",
        sub="Transformacja danych w Notebookach i na klastrach Spark",
        goal="Przekształcić dane surowe w warstwę silver i zautomatyzować proces.",
        tasks=[("2.1", "Różne sposoby pobierania danych z Lakehouse"),
               ("2.2", "Przesłanie pliku CSV z dysku i Load to Delta"),
               ("2.3–2.6", "Gotowy Notebook: z bronze do silvercleansed"),
               ("2.7–2.8", "Automatyzacja w Pipeline i kontrola wyników"),
               ("2.9", "Lakehouse gold: goldcurated")],
        outcome="trzy tabele w silvercleansed, Pipeline z ForEach i pusty goldcurated.",
        n_div="Pokaż na diagramie drogę z bronze do silver. Notebooki są gotowe, uruchamiasz je komórka po komórce. "
              "Uprzedź, że Spark session może dziś startować dłużej.",
        n_tasks="Pokaż import Notebooka i podłączenie dwóch Lakehouse. Przy 2.7 wyjaśnij ForEach i Base Parameters. "
                "Na końcu sprawdźcie razem listę tabel z zadania 2.8. Lakehouse gold będzie potrzebny w ćwiczeniu 4.",
    ),
    dict(
        img="architektura-cw2.png", big="2B", kicker="Ćwiczenie", dur="15 min", start="12:05",
        title="Ćwiczenie 2B",
        sub="Dataflow Gen2: ta sama transformacja w Power Query",
        goal="Zrobić unpivot pliku ze zniżkami bez pisania kodu.",
        tasks=[("2B.1", "Utwórz Dataflow Gen2 z poziomu Lakehouse silver"),
               ("2B.2", "Wczytaj plik NYC-Taxi-Discounts-Per-Day.csv"),
               ("2B.3", "Zrób unpivot w Power Query"),
               ("2B.4", "Zapisz wynik w Lakehouse i uruchom Dataflow Gen2"),
               ("2B.5", "Porównaj oba podejścia")],
        outcome="tabelę discounts_dataflow w silver: 6746 wierszy, zbudowaną w Power Query.",
        n_div="To krótkie ćwiczenie dla osób od Power BI i Excela. Ten sam plik CSV, ten sam wynik, inne narzędzie. "
              "Pokaż, że plik ma jedną kolumnę na każdy dzień.",
        n_tasks="Zrób unpivot na żywo w edytorze Power Query. Pokaż ustawienie miejsca docelowego na Lakehouse silver. "
                "Porównaj wynik z tabelą z Notebooka. To wstęp do slajdu „Dataflow Gen2 czy Notebook?”.",
    ),
    dict(
        img="architektura-cw3.png", big="3", kicker="Ćwiczenie", dur="40 min", start="12:20",
        title="Ćwiczenie 3",
        sub="SQL analytics endpoint, SSMS, udostępnianie i uprawnienia",
        goal="Odpytać tabele Delta w T-SQL i bezpiecznie udostępnić dane.",
        tasks=[("3.1", "Pobierz ciąg połączenia dla SQL analytics endpoint"),
               ("3.2", "Połącz się z endpointem w SSMS"),
               ("3.3", "Uruchom zapytania T-SQL na tabelach Delta"),
               ("3.4", "Udostępnij Lakehouse"),
               ("3.5", "Udostępnij Notebook do współpracy")],
        outcome="działające połączenie z SSMS i Lakehouse udostępniony tylko do odczytu.",
        n_div="Pokaż prawą stronę diagramu: konsumpcja przez SQL. Podkreśl, że SQL analytics endpoint jest tylko do odczytu. "
              "Kto nie ma SSMS, pracuje w edytorze SQL w Fabric.",
        n_tasks="Pokaż, gdzie skopiować ciąg połączenia. Przy 3.4 pokaż przycisk Share i opcję Read all with SQL analytics endpoint. "
                "Powiedz, czym różni się udostępnienie itemu od roli w workspace.",
    ),
    dict(
        img="architektura-cw4.png", big="4", kicker="Ćwiczenie", dur="80 min", start="14:00",
        title="Ćwiczenie 4",
        sub="Semantic model i raport Power BI",
        goal="Zbudować semantic model w trybie Direct Lake i raport Power BI.",
        tasks=[("4.1", "Przewidź czas przejazdu: Data Science w Lakehouse"),
               ("4.2", "Zbadaj i zwizualizuj dane w Power BI i Direct Lake"),
               ("4.3", "Opublikuj i udostępnij raport Power BI")],
        outcome="tabelę w goldcurated, semantic model w Direct Lake i własny raport.",
        n_div="Pokaż na diagramie warstwę gold i Power BI. Przypomnij, że Fabric nie tworzy już domyślnego semantic model. "
              "Ten raport każdy pokaże na koniec dnia.",
        n_tasks="Uruchom Notebook z zadania 4.1 i pokaż tabelę w gold. Zbuduj semantic model na żywo i dodaj jedną wizualizację. "
                "Wyjaśnij różnicę między Direct Lake, Import mode i DirectQuery.",
    ),
    dict(
        img="architektura-cw5.png", big="5", kicker="Ćwiczenie", dur="60 min", start="15:35",
        title="Ćwiczenie 5",
        sub="Nowości w Fabric i Data Wrangler",
        goal="Poznać nowości w Fabric i czyścić dane klikaniem w Data Wrangler.",
        tasks=[("1", "Dodaj do zakładek Fabric Monthly Updates i Fabric Blog"),
               ("2", "Runtime 1.3 i Python UDTFs"),
               ("3", "Managed Private Endpoints i Autotune Query Tuning"),
               ("4", "Spark czy Pandas"),
               ("5", "Data Wrangler i VS Code for the Web")],
        outcome="źródła nowości w zakładkach i kod wygenerowany przez Data Wrangler.",
        n_div="Ta część to mieszanka demo i samodzielnej pracy. Pokaż na diagramie dolny pasek z nowościami. "
              "Powiedz, że w produkcji używamy wersji runtime ze statusem GA.",
        n_tasks="Managed Private Endpoints pokaż jako demo. Data Wrangler uczestnicy robią sami. "
                "Pokaż, że Data Wrangler na końcu generuje kod Pandas albo PySpark.",
    ),
    dict(
        img="architektura-cw6.png", big="+", kicker="Dodatkowe", dur="własne tempo", start=None,
        title="Ćwiczenia dodatkowe",
        sub="Dla szybszych, na przerwę i na później",
        goal="Pogłębić wybrane tematy we własnym tempie, także po warsztacie.",
        tasks=[("A", "Copilot w Notebooku"),
               ("B", "Monitoring hub i Lineage"),
               ("C", "High concurrency mode i własny Spark pool"),
               ("D", "Harmonogram Notebooka"),
               ("E", "V-Order i Merge")],
        outcome="materiał do pracy po warsztacie. Wszystko jest w repozytorium.",
        n_div="Te ćwiczenia nie są obowiązkowe. Kto skończy wcześniej, wybiera temat z listy. "
              "Pokaż plik extra.md w repozytorium.",
        n_tasks="Wymień tematy i powiedz, który łączy się z którym ćwiczeniem. Copilot wymaga płatnej capacity. "
                "Zachęć do powrotu do tych zadań po warsztacie.",
    ),
]

MISTAKES = [
    dict(title="Raportowanie na brudnych danych",
         symptom="Raport czyta surowe tabele z bronze. Każdy czyści dane po swojemu, w swoim raporcie.",
         effect="Dwa raporty pokazują różne liczby. Nikt nie ufa danym.",
         fix="Czyść dane raz, w warstwie silver. Raporty buduj tylko na gold.",
         note="Zapytaj, kto widział dwa raporty z różnymi liczbami. Opowiedz własną historię z projektu. "
              "Wróć do diagramu: raport czyta tylko gold."),
    dict(title="Mieszanie warstw",
         symptom="Surowe pliki, oczyszczone tabele i agregaty leżą w jednym Lakehouse.",
         effect="Nie wiadomo, której tabeli ufać. Każda zmiana może zepsuć raport.",
         fix="Osobny Lakehouse na każdą warstwę. Dane płyną tylko z bronze do silver i do gold.",
         note="Przypomnij trzy Lakehouse z warsztatu: bronzerawdata, silvercleansed, goldcurated. "
              "Opowiedz własną historię z projektu. Podkreśl, że bronze zostaje niezmienione."),
    dict(title="Nadmiar narzędzi w jednym rozwiązaniu",
         symptom="Tę samą logikę masz w Dataflow Gen2, w Notebooku i w Pipeline.",
         effect="Trudno znaleźć błąd. Nowa osoba długo uczy się rozwiązania.",
         fix="Wybierz jedno narzędzie do transformacji. Pipeline zostaw do orkiestracji.",
         note="Nawiąż do slajdu „Dataflow Gen2 czy Notebook?”. Fabric daje wiele dróg do tego samego celu. "
              "Opowiedz własną historię z projektu."),
    dict(title="Brak konwencji nazw",
         symptom="W workspace leżą „Notebook 1”, „test2” i „final_v3”.",
         effect="Szukasz zamiast pracować. Łatwo nadpisać albo usunąć zły item.",
         fix="Ustal konwencję pierwszego dnia. Nazwa mówi, jaka to warstwa i jakie dane.",
         note="Pokaż plik naming-convention.md z repozytorium. Na warsztacie nazwa Lakehouse mówi, jaka to warstwa. "
              "Opowiedz własną historię z projektu."),
    dict(title="Zbyt szerokie uprawnienia",
         symptom="Każdy dostaje rolę Admin albo Member w workspace, bo tak jest szybciej.",
         effect="Każdy może zmienić albo usunąć dowolny item w workspace.",
         fix="Udostępnij sam Lakehouse przyciskiem Share. Zaznacz „Read all with SQL analytics endpoint”.",
         note="Nawiąż do zadania 3.4. Analityk, który tylko czyta dane, nie potrzebuje roli w workspace. "
              "Opowiedz własną historię z projektu."),
    dict(title="Brak kontroli kosztów",
         symptom="Spark session działają, choć nikt z nich nie korzysta. Capacity pracuje całą dobę.",
         effect="Capacity jest zajęta. Nowy Spark job dostaje błąd HTTP 430.",
         fix="Zmniejsz Spark pool. Anuluj zbędne Spark session w Monitoring hub. Wstrzymuj capacity z SKU F.",
         note="Nawiąż do zadania 1.4 i do kroku 13 z konfiguracji. Pokaż Monitoring hub. "
              "Opowiedz własną historię z projektu."),
]


# ---------------------------------------------------------------- slides
def build(output: Path = OUTPUT):
    d = Deck()

    # 1 -- title
    s = d.prs.slides.add_slide(d.prs.slide_layouts[L_TITLE])
    t = s.shapes.title
    place(t, 1.67, 1.55, 10.0, 1.9)
    write(t.text_frame, [[("Kupiłeś Fabric. I co teraz?", 60, NAVY, None, None)]], anchor=MSO_ANCHOR.BOTTOM)
    no_autofit(t.text_frame)
    box(s, 5.92, 3.7, 1.5, 0.08, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    sub = s.placeholders[1]
    place(sub, 1.67, 4.0, 10.0, 2.3)
    write(sub.text_frame, [
        [("Praktyczny warsztat dla osób od Power BI, SQL i Excela", 28, NAVY, None, None)],
        [("Damian Widera · dr Estera Kot", 24, NAVY, True, None)],
        [("SQLDay Lite 2026, Gdańsk, 24 września 2026", 20, GREY, None, None)],
    ], space_after=8)
    no_autofit(sub.text_frame)
    notes(s, "Przywitaj grupę i przedstaw temat. Powiedz, że to cały dzień pracy praktycznej, poziom 300, po polsku. "
             "Zapytaj, kto pracuje dziś głównie w Power BI, kto w SQL, a kto w Excelu.")

    # 2 -- full architecture
    s = d.slide(L_DARK2, "Co dziś zbudujesz", size=40, top=0.3, height=1.1)
    drop_placeholder(s, 1)
    picture(s, ARCH / "architektura.png", x=(13.333 - 10.3) / 2, y=1.52, w=10.3)
    notes(s, "Przejdź diagram od lewej do prawej: źródła danych, ładowanie, OneLake z trzema Lakehouse, konsumpcja. "
             "Powiedz, że ten sam diagram wraca przed każdym ćwiczeniem z podświetlonym fragmentem. "
             "Na końcu dnia każdy ma raport Power BI w trybie Direct Lake.")

    # 3 -- speakers
    s = d.slide(L_TITLE_ONLY, "Prowadzący")
    people = [("Damian Widera", "Data Solution Architect", "MVP Data Platform"),
              ("dr Estera Kot", "CTO", "Clouds on Mars")]
    for i, (name, role1, role2) in enumerate(people):
        x = LEFT + i * 5.9
        box(s, x, 1.85, 5.6, 4.9, fill=MIST, radius=0.05)
        tbc(s, x + 0.35, 2.2, 2.2, 2.75, "⟦TBC: zdjęcie⟧")
        textbox(s, x + 2.75, 2.2, 2.8, 2.75, [
            [(name, 32, NAVY, True)],
            [(role1, 22, NAVY, False)],
            [(role2, 22, GREY, False)],
        ], space_after=6)
        box(s, x + 0.35, 5.3, 0.7, 0.07, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
        tbc(s, x + 0.35, 5.6, 4.9, 0.8, "⟦TBC: jedno zdanie o prowadzącym⟧", size=16)
    notes(s, "Każdy prowadzący przedstawia się sam, w dwóch zdaniach. Powiedz, z czym można do kogo przyjść. "
             "Oboje pomagamy przy stolikach przez cały dzień.")

    # 4 -- audience
    s = d.slide(L_TITLE_ONLY, "Dla kogo jest ten warsztat")
    who = ["analitycy Power BI", "osoby pracujące z SQL", "użytkownicy Excela i Power Query",
           "deweloperzy BI", "początkujący inżynierzy danych", "konsultanci, którzy porządkują podstawy Fabric"]
    cw, ch, gap = 3.7, 1.75, 0.2
    for i, label in enumerate(who):
        x = LEFT + (i % 3) * (cw + gap)
        y = 1.85 + (i // 3) * (ch + gap)
        c = box(s, x, y, cw, ch, fill=NAVY, radius=0.08)
        box(s, x + 0.25, y + 0.22, 0.6, 0.07, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
        textbox(s, x + 0.17, y + 0.38, cw - 0.34, ch - 0.45, [[(label, 24, WHITE, False)]])
    b = box(s, LEFT, 5.75, WIDTH, 1.0, fill=MIST, radius=0.08)
    write(b.text_frame, [
        [("Nie musisz znać Sparka, Pythona, DevOps ani Gita. ", 24, NAVY, True),
         ("Notebooki są gotowe.", 24, NAVY, False)]],
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0.3, 0.05, 0.3, 0.05))
    notes(s, "Warsztat jest dla osób, które zaczynają z Fabric albo znają go tylko od strony Power BI. "
             "Uspokój grupę: Notebooki są gotowe, uruchamiasz je komórka po komórce. Tłumaczymy, co się dzieje w każdym kroku.")

    # 5 -- agenda
    s = d.slide(L_TITLE_ONLY, "Agenda")
    left = [("08:30", "Wprowadzenie i konfiguracja", False),
            ("09:15", "Ćwiczenie 1: Pipeline i Shortcut", False),
            ("10:35", "Przerwa kawowa", True),
            ("10:50", "Ćwiczenie 2 i 2B: Notebook, Dataflow Gen2", False),
            ("12:20", "Ćwiczenie 3: SQL analytics endpoint", False)]
    right = [("13:00", "Lunch", True),
             ("14:00", "Ćwiczenie 4: semantic model i raport", False),
             ("15:20", "Przerwa", True),
             ("15:35", "Ćwiczenie 5: nowości i Data Wrangler", False),
             ("16:35", "Finał: raporty, błędy, ankieta", False)]
    for col, rows in enumerate((left, right)):
        x = LEFT + col * 5.85
        for i, (tm, label, is_break) in enumerate(rows):
            y = 1.8 + i * 0.88
            box(s, x, y, 5.65, 0.74, fill=MIST, radius=0.12)
            chip = box(s, x, y, 1.15, 0.74, fill=LORANGE if is_break else NAVY, radius=0.12)
            write(chip.text_frame, [[(tm, 22, NAVY if is_break else WHITE, True)]],
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))
            textbox(s, x + 1.25, y, 4.35, 0.74, [[(label, 22, GREY if is_break else NAVY, False)]],
                    anchor=MSO_ANCHOR.MIDDLE)
    textbox(s, LEFT, 6.3, WIDTH, 0.6,
            [[("Godziny mogą się przesunąć. ", 20, NAVY, True),
              ("Dopasujemy je do tempa grupy i harmonogramu organizatora.", 20, NAVY, False)]])
    notes(s, "Przejdź agendę w minutę. Powiedz, że godziny dopasujemy do tempa grupy i do harmonogramu organizatora. "
             "Przerwy są propozycją. Ćwiczenie 2B o Dataflow Gen2 zaczynamy o 12:05, zaraz po ćwiczeniu 2. ⟦TBC: godziny po publikacji harmonogramu SQLDay Lite⟧")

    # 6 -- rules
    s = d.slide(L_DARK, "Zasady pracy")
    drop_placeholder(s, 1)
    rules = [("Własne tempo", "Przerwy są propozycją."),
             ("Pytania mile widziane", "Pytaj w dowolnej chwili."),
             ("Nazwy funkcji po angielsku", "Tak jak w interfejsie."),
             ("Tryb incognito", "Bez logowania do firmowego tenanta."),
             ("Konwencja nazw", "Od niej zależy, czy ćwiczenia zadziałają.")]
    cw, ch, gap = 3.7, 2.35, 0.2
    for i, (h, b) in enumerate(rules):
        x = LEFT + (i % 3) * (cw + gap)
        y = 1.85 + (i // 3) * (ch + gap)
        card(s, x, y, cw, ch, h, b, dark=True, head_size=26, body_size=22, number=f"{i + 1}")
    tbc(s, LEFT + 2 * (cw + gap), 1.85 + ch + gap, cw, ch, "⟦TBC: link do grupy uczestników⟧", dark=True, size=20)
    notes(s, "Omów pięć zasad. Nazwy funkcji, przycisków i opcji zostają po angielsku, bo tak wyglądają na ekranie. "
             "Tryb incognito chroni przed logowaniem do firmowego tenanta. Konwencję nazw wystarczy przeczytać i stosować.")

    # 7 -- five concepts
    s = d.slide(L_TITLE_ONLY, "Fabric w 5 pojęciach")
    terms = [("tenant", "Twoja organizacja w chmurze Microsoft."),
             ("capacity", "Pula mocy obliczeniowej, którą kupuje firma."),
             ("workspace", "Kontener na elementy projektu. Tu nadajesz role."),
             ("OneLake", "Jeden logiczny data lake dla całego tenanta."),
             ("item", "Każdy obiekt w workspace, np. Lakehouse, Notebook, raport.")]
    for i, (h, b) in enumerate(terms):
        x = LEFT + (i % 3) * (cw + gap)
        y = 1.85 + (i // 3) * (ch + gap)
        card(s, x, y, cw, ch, h, b, dark=True, head_size=32, body_size=22)
    g = box(s, LEFT + 2 * (cw + gap), 1.85 + ch + gap, cw, ch, fill=MIST, radius=0.06)
    write(g.text_frame, [[("Pełny słowniczek", 26, NAVY, True)],
                         [("35 pojęć w README repozytorium, z ID od G01 do G35.", 22, NAVY, False)]],
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0.25, 0.1, 0.25, 0.1), space_after=6)
    notes(s, "Wyjaśnij pięć pojęć na przykładzie dzisiejszego środowiska. Jeden tenant ma jeden OneLake. "
             "Wszystko, co uruchamiasz, zużywa capacity przypisaną do workspace. Pokaż te miejsca w portalu Fabric.")

    # 8 -- medallion
    s = d.slide(L_DARK_WARM, "Medallion architecture")
    drop_placeholder(s, 1)
    layers = [("bronze", "dane surowe", "bronzerawdata", BRONZE),
              ("silver", "dane oczyszczone", "silvercleansed", SILVER),
              ("gold", "dane gotowe do raportowania", "goldcurated", LORANGE)]
    cw3, gap3 = 3.3, 0.8
    for i, (name, desc, lh, accent) in enumerate(layers):
        x = LEFT + i * (cw3 + gap3)
        box(s, x, 1.9, cw3, 3.7, fill=NAVY, line=accent, line_w=1.5, radius=0.06)
        box(s, x + 0.3, 2.2, cw3 - 0.6, 0.09, fill=accent, shape=MSO_SHAPE.RECTANGLE)
        textbox(s, x + 0.22, 2.45, cw3 - 0.44, 3.0, [
            [(name, 44, WHITE, True)],
            [(desc, 24, WHITE, False)],
            [(lh, 24, LORANGE, False, "Consolas")],
        ], space_after=10)
        if i < 2:
            box(s, x + cw3 + 0.15, 3.5, gap3 - 0.3, 0.5, fill=ORANGE, shape=MSO_SHAPE.RIGHT_ARROW)
    textbox(s, LEFT, 6.0, WIDTH, 0.6,
            [[("Raporty buduj tylko na gold. ", 24, LORANGE, True),
              ("Każda warstwa ma na warsztacie własny Lakehouse.", 24, WHITE, False)]])
    notes(s, "Wyjaśnij trzy warstwy i powiedz, że na warsztacie każda ma własny Lakehouse. "
             "Dane w bronze zostają niezmienione. Raporty budujemy tylko na gold. To pierwsza obrona przed bałaganem.")

    # 9 -- exercises: divider + goals
    for ex in EXERCISES:
        s = d.slide(L_DARK_WARM, ex["title"], size=38, top=0.25, height=1.25, width=9.4,
                    sub=ex["sub"], sub_color=LORANGE, sub_size=22)
        drop_placeholder(s, 1)
        picture(s, ARCH / ex["img"], x=0.45, y=1.62, w=10.0)
        sx, sw = 10.7, 2.2
        box(s, sx, 1.62, sw, 5.625, fill=NAVY, line=ORANGE, line_w=1.25, radius=0.06)
        textbox(s, sx, 2.45, sw, 0.45, [[(ex["kicker"].upper(), 16, LORANGE, True)]], align=PP_ALIGN.CENTER)
        textbox(s, sx, 2.85, sw, 1.7, [[(ex["big"], 96 if len(ex["big"]) == 1 else 80, ORANGE, True)]],
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        box(s, sx + 0.75, 4.8, 0.7, 0.07, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
        rows = [[(ex["dur"], 28 if len(ex["dur"]) < 8 else 22, WHITE, True)]]
        if ex["start"]:
            rows.append([("start " + ex["start"], 20, LORANGE, False)])
        textbox(s, sx, 5.05, sw, 1.4, rows, align=PP_ALIGN.CENTER, space_after=6)
        notes(s, ex["n_div"])

        s = d.slide(L_LIGHT, ex["title"] + ": cel i zadania")
        drop_placeholder(s, 1)
        textbox(s, LEFT, 1.62, 10.6, 0.6, [[("Cel: ", 22, ORANGE, True), (ex["goal"], 22, NAVY, False)]])
        for i, (tid, label) in enumerate(ex["tasks"]):
            y = 2.35 + i * 0.71
            box(s, LEFT, y, 7.6, 0.62, fill=MIST, radius=0.12)
            chip = box(s, LEFT, y, 1.2, 0.62, fill=NAVY, radius=0.12)
            write(chip.text_frame, [[(tid, 20, WHITE, True)]], align=PP_ALIGN.CENTER,
                  anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0))
            textbox(s, LEFT + 1.3, y, 6.25, 0.62, [[(label, 20, NAVY, False)]], anchor=MSO_ANCHOR.MIDDLE)
        picture(s, ARCH / ex["img"], x=8.72, y=2.35, w=3.7)
        b = box(s, LEFT, 6.0, WIDTH, 0.9, fill=NAVY, radius=0.1)
        write(b.text_frame, [[("Po tym ćwiczeniu masz: ", 22, LORANGE, True), (ex["outcome"], 22, WHITE, False)]],
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0.3, 0.04, 0.3, 0.04))
        notes(s, ex["n_tasks"])

    # 10 -- design decisions
    def compare(title, lh, lb, rh, rb, verdict, note):
        s = d.slide(L_COMPARE, title)
        geo = {1: (LEFT, 1.75, 5.55, 0.8), 2: (LEFT, 2.75, 5.55, 2.9),
               3: (6.87, 1.75, 5.55, 0.8), 4: (6.87, 2.75, 5.55, 2.9)}
        for idx, (x, y, w, h) in geo.items():
            ph = s.placeholders[idx]
            ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
        for idx, head in ((1, lh), (3, rh)):
            write(s.placeholders[idx].text_frame, [[(head, 36, NAVY, True, None)]], anchor=MSO_ANCHOR.BOTTOM)
            no_autofit(s.placeholders[idx].text_frame)
        for idx, body in ((2, lb), (4, rb)):
            write(s.placeholders[idx].text_frame, [[(b_, 26, NAVY, None, None)] for b_ in body], space_after=10)
            no_autofit(s.placeholders[idx].text_frame)
        for x in (LEFT, 6.87):
            box(s, x + 0.1, 2.58, 1.2, 0.07, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
        b = box(s, LEFT, 5.85, WIDTH, 1.0, fill=NAVY, radius=0.1)
        write(b.text_frame, [verdict], align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
              margins=(0.3, 0.04, 0.3, 0.04))
        notes(s, note)
        return s

    compare("Lakehouse czy Warehouse?",
            "Lakehouse", ["pliki i tabele w jednym miejscu", "Spark i Python",
                          "SQL tylko do odczytu, przez SQL analytics endpoint"],
            "Warehouse", ["pełny T-SQL: INSERT, UPDATE, DELETE", "procedury", "zespół pracuje w SQL"],
            [("Oba zapisują Delta w OneLake, ", 24, LORANGE, True), ("więc możesz je łączyć.", 24, WHITE, False)],
            "Zapytaj grupę, kto pracuje głównie w T-SQL. Lakehouse wystarcza na start, tak jak na warsztacie. "
            "Warehouse wybierasz, gdy potrzebujesz zapisu w T-SQL i procedur. To nie jest wybór na zawsze, bo oba zapisują Delta w OneLake.")
    compare("Dataflow Gen2 czy Notebook?",
            "Dataflow Gen2", ["znasz Power Query", "logikę klikasz, bez kodu", "dane małe i średnie"],
            "Notebook", ["duże dane", "złożona logika", "kod w repozytorium i testy"],
            [("Wybierz jedno narzędzie do transformacji ", 24, LORANGE, True),
             ("i trzymaj się go w całym rozwiązaniu.", 24, WHITE, False)],
            "Nawiąż do ćwiczenia 2 i 2B: ta sama transformacja w dwóch narzędziach. "
            "Dataflow Gen2 to naturalny wybór dla osób od Power BI i Excela. Notebook wybierasz do dużych danych i złożonej logiki.")

    s = d.slide(L_TWO, "Jak dostarczyć dane?")
    ways = [("Pipeline z Copy activity", "cykliczne kopiowanie danych"),
            ("Shortcut", "dane zostają u źródła"),
            ("Upload", "mały plik, jednorazowo"),
            ("Dataflow Gen2", "pobranie danych z transformacją")]
    for idx, x in ((1, LEFT), (2, 6.87)):
        ph = s.placeholders[idx]
        ph.left, ph.top, ph.width, ph.height = Inches(x + 0.25), Inches(1.9), Inches(5.3), Inches(3.8)
        pars = []
        for h, b_ in ways[(idx - 1) * 2:(idx - 1) * 2 + 2]:
            pars.append([(h, 32, NAVY, True, None)])
            pars.append([(b_, 24, NAVY, None, None)])
        write(ph.text_frame, pars, bullets=False, space_after=4)
        ph.text_frame.paragraphs[1].space_after = Pt(44)
        no_autofit(ph.text_frame)
        box(s, x, 2.0, 0.09, 1.25, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
        box(s, x, 3.75, 0.09, 1.25, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    b = box(s, LEFT, 5.85, WIDTH, 1.0, fill=NAVY, radius=0.1)
    write(b.text_frame, [[("Trzy z tych dróg znasz już z ćwiczeń 1 i 2. ", 24, LORANGE, True),
                          ("Czwartą poznajesz w ćwiczeniu 2B.", 24, WHITE, False)]],
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margins=(0.3, 0.04, 0.3, 0.04))
    notes(s, "Połącz każdą drogę z ćwiczeniem: Pipeline z zadania 1.1, Shortcut z 1.3, Upload z 2.2, Dataflow Gen2 z 2B. "
             "Zapytaj, gdzie dziś leżą dane uczestników. Shortcut nie kopiuje danych, widzisz je u siebie jak własne.")

    # 11 -- leave for later
    s = d.slide(L_DARK_BLUE, "Zostaw na później")
    drop_placeholder(s, 1)
    later = [("Managed Private Endpoints", "Tylko gdy źródło stoi za firewallem."),
             ("custom Spark pool", "Starter pool startuje w kilka sekund."),
             ("Real-Time Intelligence", "Najpierw dane wsadowe, potem strumienie."),
             ("Data Science", "Modele dopiero na czystych danych."),
             ("deployment pipelines", "Najpierw jeden porządny workspace.")]
    for i, (h, b_) in enumerate(later):
        y = 1.8 + i * 1.0
        box(s, LEFT, y, WIDTH, 0.85, fill=NAVY, radius=0.12)
        chip = box(s, LEFT, y, 4.5, 0.85, fill=ORANGE, radius=0.12)
        write(chip.text_frame, [[(h, 24, NAVY, True)]], align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
              margins=(0.25, 0, 0.1, 0))
        textbox(s, LEFT + 4.7, y, WIDTH - 4.9, 0.85, [[(b_, 24, WHITE, False)]], anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Powiedz, że te funkcje są dobre, ale nie są potrzebne w pierwszym projekcie. "
             "Managed Private Endpoints wyłączają Starter pool w workspace. Custom pool startuje kilka minut. "
             "Wróć do nich, gdy pierwszy przepływ danych działa i ma właściciela.")

    # 12 -- mistakes overview + one slide per mistake
    s = d.slide(L_DARK2, "Najczęstsze błędy początkujących")
    drop_placeholder(s, 1)
    cw, ch, gap = 3.7, 2.3, 0.2
    for i, m in enumerate(MISTAKES):
        x = LEFT + (i % 3) * (cw + gap)
        y = 1.9 + (i // 3) * (ch + gap)
        box(s, x, y, cw, ch, fill=NAVY, line=ORANGE, line_w=1.0, radius=0.07)
        textbox(s, x + 0.2, y + 0.12, 0.9, ch - 0.24, [[(str(i + 1), 60, ORANGE, True)]], anchor=MSO_ANCHOR.MIDDLE)
        textbox(s, x + 1.05, y + 0.12, cw - 1.2, ch - 0.24, [[(m["title"], 26, WHITE, False)]],
                anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Ten blok zamyka dzień. To błędy, które widzimy w pierwszych projektach Fabric. "
             "Zapytaj, który z nich grupa już zna z własnej pracy. Każdy omawiamy na osobnym slajdzie.")

    for i, m in enumerate(MISTAKES):
        s = d.slide(L_TITLE_ONLY, f"Błąd {i + 1}: {m['title'][0].lower() + m['title'][1:]}", size=36)
        cols = [("Objaw", m["symptom"], False), ("Skutek", m["effect"], False), ("Jak tego uniknąć", m["fix"], True)]
        cw, gap = 3.7, 0.2
        for j, (h, b_, good) in enumerate(cols):
            x = LEFT + j * (cw + gap)
            card(s, x, 1.85, cw, 3.75, h, b_, dark=True, head_size=30, body_size=24,
                 accent=LORANGE if good else ORANGE)
        tbc(s, LEFT, 5.85, WIDTH, 0.9, "⟦TBC: historia z projektu⟧", size=20)
        notes(s, m["note"])

    # 13 -- show your report
    s = d.slide(L_BLUE, "Pokaż swój raport")
    ph = s.placeholders[1]
    place(ph, LEFT, 2.1, 10.5, 4.2)
    write(ph.text_frame, [
        [("Pokaż swój raport Power BI z ćwiczenia 4", 34, WHITE, None, None)],
        [("Powiedz, co cię dziś zaskoczyło", 34, WHITE, None, None)],
        [("Powiedz, co zrobisz w Fabric jako pierwsze po powrocie do pracy", 34, WHITE, None, None)],
    ], space_after=20)
    no_autofit(ph.text_frame)
    notes(s, "Każdy uczestnik pokazuje swój raport Power BI, krótko, przy swoim stoliku albo na rzutniku. "
             "Nie oceniamy wyglądu. Pytamy, co zaskoczyło i co uczestnik zrobi jako pierwsze po powrocie do pracy.")

    # 14 -- survey
    s = d.slide(L_TITLE_ONLY, "Twoja opinia")
    textbox(s, LEFT, 2.3, 5.9, 3.6, [
        [("Zeskanuj kod QR albo otwórz formularz.", 32, NAVY, False)],
        [("forms.cloud.microsoft/e/hXHYaB8pDb", 24, NAVY, True)],
        [("Twoje odpowiedzi pomogą nam poprawić kolejną edycję.", 24, GREY, False)],
    ], space_after=18)
    box(s, LEFT + 0.05, 2.12, 0.7, 0.07, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    picture(s, QR, x=7.0, y=1.6, w=5.4)
    notes(s, "Zostaw ten slajd na ekranie na kilka minut. Poproś o wypełnienie ankiety teraz, zanim grupa się rozejdzie. "
             "Odpowiedzi pomogą poprawić kolejną edycję warsztatu.")

    # 15 -- what next
    s = d.slide(L_DARK, "Co dalej")
    drop_placeholder(s, 1)
    nxt = [("Repozytorium", "github.com/DamianWidera/SQLDayLite2026"),
           ("Ćwiczenia dodatkowe", "Copilot, Lineage, Merge i inne. We własnym tempie."),
           ("Ściąga PDF", "Znajdziesz ją w repozytorium."),
           ("Fabric Monthly Updates", "Nowości z Fabric w jednym miejscu.")]
    cw, ch, gap = 5.65, 2.35, 0.2
    for i, (h, b_) in enumerate(nxt):
        x = LEFT + (i % 2) * (cw + gap)
        y = 1.85 + (i // 2) * (ch + gap)
        card(s, x, y, cw, ch, h, b_, dark=True, head_size=30, body_size=22)
    notes(s, "Repozytorium zostaje dostępne po warsztacie, razem z instrukcjami i Notebookami. "
             "Pokaż, gdzie leży ściąga PDF i ćwiczenia dodatkowe. Poleć Fabric Monthly Updates jako jedno źródło nowości.")

    # 16 -- thank you
    s = d.prs.slides.add_slide(d.prs.slide_layouts[L_SECTION])
    t = s.shapes.title
    place(t, LEFT, 1.6, 9.0, 1.9)
    write(t.text_frame, [[("Dziękujemy", 66, NAVY, None, None)]], anchor=MSO_ANCHOR.BOTTOM)
    no_autofit(t.text_frame)
    box(s, 1.0, 3.62, 1.5, 0.08, fill=ORANGE, shape=MSO_SHAPE.RECTANGLE)
    drop_placeholder(s, 1)
    for i, name in enumerate(("Damian Widera", "dr Estera Kot")):
        x = LEFT + i * 4.6
        textbox(s, x, 3.95, 4.3, 0.6, [[(name, 28, NAVY, True)]])
        tbc(s, x + 0.05, 4.65, 4.0, 0.7, "⟦TBC: kontakt⟧", size=18)
    picture(s, LOGO_DC, x=9.9, y=6.05, w=2.5)
    notes(s, "Podziękuj za cały dzień pracy. Powiedz, gdzie można nas znaleźć i jak zadać pytanie po warsztacie. "
             "Przypomnij o ankiecie i o konferencji SQLDay Lite następnego dnia.")

    d.save(output)
    return len(d.prs.slides)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT
    n = build(out)
    print(f"Saved {out.name}: {n} slides")
